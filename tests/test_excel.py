from datetime import timedelta

from openpyxl import load_workbook

from gia.excel import COL, FIRST_DATA_ROW, HEADER_ROW, build_workbook
from gia.models import OrgType, Status
from tests.conftest import NOW, make_posting, make_source


def _postings():
    return [
        make_posting(
            "수영강사 모집", org="창원시설공단", field="sports", region=["창원"],
            deadline=NOW + timedelta(days=2), employment_type="시간강사", pay="시급 3만원",
            one_line_summary="창원 체육센터 수영강사", qualifications=["생활체육지도사 2급"],
            apply_method="이메일", status=Status.closing_soon,
        ),
        make_posting(
            "시민강좌 강사 모집", org="김해시", org_type=OrgType.local_gov, field="lifelong",
            region=["김해"], deadline=NOW + timedelta(days=20), status=Status.new, flags=["재공고"],
        ),
        make_posting("지난 공고", org="밀양시", org_type=OrgType.local_gov, field="school",
                     deadline=NOW - timedelta(days=1), status=Status.expired),
    ]


def _sources():
    a = make_source("gimhae", "local_gov", type="html_list", list_url="https://x.example/list")
    b = make_source("todo_src", "local_public", type="html_list", list_url="TODO")
    return [a, b]


def test_workbook_sheets_and_rows(tmp_path):
    path = build_workbook(_postings(), _sources(), tmp_path / "out.xlsx", NOW,
                          source_names={"src": "테스트 소스"})
    wb = load_workbook(path)
    assert wb.sheetnames == ["공고목록", "요약", "수집원 현황"]

    ws = wb["공고목록"]
    # 마감 가까운 순, 마감된 공고는 뒤로
    titles = [ws[f'{COL["공고 제목"]}{r}'].value for r in range(FIRST_DATA_ROW, FIRST_DATA_ROW + 3)]
    assert titles == ["수영강사 모집", "시민강좌 강사 모집", "지난 공고"]
    assert ws[f'{COL["상태"]}{FIRST_DATA_ROW}'].value == "마감임박"
    assert ws[f'{COL["분야"]}{FIRST_DATA_ROW}'].value == "체육"
    assert ws[f'{COL["강사료"]}{FIRST_DATA_ROW}'].value == "시급 3만원"
    assert ws[f'{COL["자격요건"]}{FIRST_DATA_ROW}'].value == "생활체육지도사 2급"
    assert ws[f'{COL["비고"]}{FIRST_DATA_ROW + 1}'].value == "재공고"
    assert ws.freeze_panes == f"A{FIRST_DATA_ROW}"
    assert ws.auto_filter.ref.startswith(f"A{HEADER_ROW}:")


def test_deadline_is_timezone_naive_datetime(tmp_path):
    """엑셀은 시간대를 다루지 못하므로 tz 를 떼고 KST 벽시계 값으로 써야 한다."""
    path = build_workbook(_postings(), _sources(), tmp_path / "out.xlsx", NOW)
    ws = load_workbook(path)["공고목록"]
    dl = ws[f'{COL["마감일"]}{FIRST_DATA_ROW}'].value
    assert dl.tzinfo is None
    assert dl.strftime("%Y-%m-%d") == (NOW + timedelta(days=2)).strftime("%Y-%m-%d")
    assert ws[f'{COL["마감일"]}{FIRST_DATA_ROW}'].number_format == "yyyy-mm-dd"


def test_dday_and_summary_use_formulas(tmp_path):
    """D-day 와 집계는 값이 아니라 수식이어야 파일을 열 때 갱신된다."""
    path = build_workbook(_postings(), _sources(), tmp_path / "out.xlsx", NOW)
    wb = load_workbook(path)
    dday = wb["공고목록"][f'{COL["D-day"]}{FIRST_DATA_ROW}'].value
    assert dday.startswith("=IF(") and "TODAY()" in dday

    summary = wb["요약"]
    formulas = [c.value for row in summary.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert any(f.startswith("=COUNTIF(") for f in formulas)
    assert any(f.startswith("=COUNTIFS(") for f in formulas)
    assert any(f.startswith("=SUM(") for f in formulas)
    # 집계는 공고목록 시트를 참조한다 (하드코딩 금지)
    assert all("'공고목록'!" in f for f in formulas if f.startswith(("=COUNTIF(", "=COUNTIFS(")))


def test_hyperlink_on_source_url(tmp_path):
    path = build_workbook(_postings(), _sources(), tmp_path / "out.xlsx", NOW)
    ws = load_workbook(path)["공고목록"]
    cell = ws[f'{COL["원문 링크"]}{FIRST_DATA_ROW}']
    assert cell.value == "원문 보기"
    assert cell.hyperlink.target.startswith("https://example.org/")


def test_sources_sheet_flags_unconfigured(tmp_path):
    path = build_workbook(_postings(), _sources(), tmp_path / "out.xlsx", NOW)
    ws = load_workbook(path)["수집원 현황"]
    rows = {ws.cell(row=r, column=2).value: [ws.cell(row=r, column=c).value for c in range(1, 10)]
            for r in range(FIRST_DATA_ROW, FIRST_DATA_ROW + 2)}
    assert "미설정" in rows["todo_src"][8]
    assert rows["gimhae"][8] in (None, "")


def test_empty_store_still_builds(tmp_path):
    path = build_workbook([], _sources(), tmp_path / "empty.xlsx", NOW)
    wb = load_workbook(path)
    assert wb["공고목록"].auto_filter.ref is not None
    assert wb["요약"]["A4"].value == "상태별"
