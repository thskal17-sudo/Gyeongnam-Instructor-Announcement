"""강사잇다 양식 채우기와 그 재료(본문 발췌·백필·항목 추출)."""
from __future__ import annotations

from datetime import datetime, timedelta
from io import BytesIO

import httpx
import pytest
from openpyxl import load_workbook

from gia.extract.deadline import KST
from gia.extract.detail import (
    clean_detail, clean_title, expand_school_name, extract_documents,
    extract_headcount, extract_qualifications, extract_schedule,
)
from gia.models import OrgType, RawPosting, SourceRef, Status
from gia.report.gangsaitda import COLS, TEMPLATE, build_gangsaitda, region_label
from gia.store import Store
from tests.conftest import NOW, make_posting, make_source

# 2026-09-28 러너 덤프에서 옮긴 경남교육청 구인구직포털 본문 (div.bd-view__vcontent div.txt)
MUJIGAE = """2026학년도 무지개초등학교 돌봄교실  특기·적성 프로그램 개인위탁 외부강사 모집 공고
1. 모집강좌: 클레이(1명),영어(1명),책놀이(1명)
2.계약기간: 2026년 10월 16일 ~ 2027년 2월 26일
(학사일정에 따라 변동될 수 있음, 예산 소진 시 미 운영함)
3. 모집 세부 사항
가. 공고 기간：2026. 9. 28.(월) ∼ 2026. 10. 2.(금)
나. 서류 접수：2026. 9. 28.(월) ∼ 2026. 10. 2.(금) 12:00까지
다. 접수 장소：전자 메일 접수(taechan@korea.kr)
사. 문의: 무지개초 돌봄4교실 전화 070-5091-7588(10:30~18:30)
※세부사항은 첨부파일을 참고 바랍니다."""
GYEONGHAE = """2026학년도 경해여자고등학교에 근무할 결원대체강사(전일제 강사)를 다음과 같이 채용 공고합니다.
1. 채용과목 및 응모 자격
가. 채용과목 : 국어
나. 채용인원 : 1명
다. 응모자격 : 해당 교과목 교사 자격증 소지자(단, 남자는 병역을 필하였거나 면제자)
2. 임용 기간 : 2026. 10. 12. ~ 2026. 10. 22.(주 16차시)
3. 제출서류
◦ 응시원서 1부(붙임서식, 사진첨부, 연락 가능한 전화번호 반드시 기재)
◦ 자기소개서 1부(붙임서식)
4. 접수기간 및 접수처
가. 접수기간 : 2026. 09. 22.(화) ～ 2026. 09. 28.(월) 11:00까지
나. 접수방법 : 학교 이메일 접수(gyeonghae@hanmail.net)
5. 전형 방법 및 일시
마. 기타 자세한 사항은 경해여자고등학교(☎055-746-4183)로 문의하시기 바랍니다."""
DAEJIN = """2026. 대진초 집중돌봄교실 특기적성 프로그램 외부 강사 모집 재공고
-접수기간: 2026. 9. 21.(월) ~ 2026. 9. 28.(월)
-접수방법: 대진초 교무실 직접 접수 또는 전자메일 접수
*자세한 사항은 첨부물을 참조 하시기 바랍니다."""


# ---- 항목 추출 -------------------------------------------------------

def test_extract_from_real_school_postings():
    assert extract_schedule(MUJIGAE) == "2026년 10월 16일 ~ 2027년 2월 26일"
    assert extract_headcount(MUJIGAE) == 3  # 클레이·영어·책놀이 각 1명

    assert extract_schedule(GYEONGHAE) == "2026. 10. 12. ~ 2026. 10. 22.(주 16차시)"
    assert extract_headcount(GYEONGHAE) == 1
    assert extract_qualifications(GYEONGHAE) == ["해당 교과목 교사 자격증 소지자(단, 남자는 병역을 필하였거나 면제자)"]
    assert extract_documents(GYEONGHAE).startswith("응시원서 1부")
    assert "자기소개서 1부" in extract_documents(GYEONGHAE)


def test_extract_leaves_blank_when_body_points_to_attachment():
    """'첨부 참조'만 있는 공고에서 일정을 지어내지 않는다."""
    assert extract_schedule(DAEJIN) == ""
    assert extract_headcount(DAEJIN) is None


def test_clean_detail_drops_board_chrome():
    text = "강사 모집\n목록\n미리보기\n1. 모집분야: 요가\n[전화번호]\n무지개초 문의처:"
    assert clean_detail(text, title="강사 모집") == "1. 모집분야: 요가"


@pytest.mark.parametrize("raw,expected", [
    ("2026. 대진초 집중돌봄교실 특기적성 외부 강사 모집", "대진초 집중돌봄교실 특기적성 외부 강사 모집"),
    ("2026년 한일여고 방과후학교(금융) 외부강사 모집공고", "한일여고 방과후학교(금융) 외부강사 모집공고"),
    ("2026학년도 가람초 계약제교원 채용 공고", "가람초 계약제교원 채용 공고"),
    ("2026. 2학기 돌봄교실 특기적성(~10.2.) 강사 모집", "돌봄교실 특기적성 강사 모집"),
    ("경해여자고등학교 결원대체강사(국어) 채용 공고(4차)", "경해여자고등학교 결원대체강사(국어) 채용 공고(4차)"),
])
def test_clean_title_strips_dates_only(raw, expected):
    assert clean_title(raw) == expected


@pytest.mark.parametrize("org,expected", [
    ("경해여자고", "경해여자고등학교"), ("한일여고", "한일여자고등학교"), ("대진초", "대진초등학교"),
    ("거창중", "거창중학교"), ("양산교육지원청", "양산교육지원청"), ("금송중학교", "금송중학교"),
])
def test_expand_school_name(org, expected):
    assert expand_school_name(org) == expected


def test_region_label():
    assert region_label(["창원"]) == "경남 창원시"
    assert region_label(["하동"]) == "경남 하동군"
    assert region_label(["경남"]) == "경남"
    assert region_label([]) == "경남"


# ---- 양식 채우기 ------------------------------------------------------

def _store(tmp_path, *postings) -> Store:
    st = Store(tmp_path)
    for p in postings:
        st.upsert(p)
    return st


def test_build_fills_template_and_holds_incomplete_rows(tmp_path):
    later = NOW + timedelta(days=10)
    full = make_posting("2026. 무지개초 돌봄교실 외부강사 모집", org="무지개초", deadline=later,
                        org_type=OrgType.edu_office, region=["진주"], body_excerpt=MUJIGAE,
                        status=Status.active)
    no_body = make_posting("대진초 집중돌봄교실 외부 강사 모집", org="대진초", deadline=later - timedelta(days=5),
                           org_type=OrgType.edu_office, region=["김해"], body_excerpt=DAEJIN, status=Status.active)
    doubtful = make_posting("청소년방과후아카데미 기간제근로자 채용", org="하동군", deadline=later,
                            region=["하동"], body_excerpt=GYEONGHAE, flags=["판별유보"], relevance_score=60,
                            status=Status.active)
    expired = make_posting("지난 공고", deadline=NOW - timedelta(days=1), body_excerpt=MUJIGAE, status=Status.expired)

    blob, ready, held = build_gangsaitda(_store(tmp_path, full, no_body, doubtful, expired), NOW)
    assert (ready, held) == (1, 2)

    wb = load_workbook(BytesIO(blob))
    assert wb.sheetnames == ["공고", "예시", "안내"]  # 예시·안내 시트는 그대로
    assert wb["예시"]["A2"].value == "방과후 로봇과학 강사 모집"
    ws = wb["공고"]
    assert [c.value for c in ws[1]][:13] == COLS
    assert ws["N1"].value == "처리" and ws["O1"].value == "메모"

    row = {h: ws.cell(2, i + 1).value for i, h in enumerate(COLS + ["처리", "메모"])}
    assert row["제목"] == "무지개초 돌봄교실 외부강사 모집"  # 연도 머리 뺌
    assert row["기관명"] == "무지개초등학교"
    assert row["지역"] == "경남 진주시"
    assert row["마감일"] == later.date().isoformat()  # 문자열, 양식 예시와 같은 꼴
    assert row["수업 일정"] == "2026년 10월 16일 ~ 2027년 2월 26일"
    assert "1. 모집강좌: 클레이(1명)" in row["상세 내용"]
    assert row["모집 인원"] == 3
    assert row["지원 이메일"] is None  # 가리지 않은 메일을 공개 저장소에 쌓지 않으려고 비워 둔다
    assert row["원문 링크"].startswith("https://")
    assert row["처리"] is None
    assert ws.cell(2, 6).alignment.wrap_text  # 칸 안 줄바꿈이 보이게

    held_rows = {ws.cell(r, 1).value: (ws.cell(r, 14).value, ws.cell(r, 15).value) for r in (3, 4)}
    assert held_rows["대진초 집중돌봄교실 외부 강사 모집"][0] == "보류"
    assert "수업 일정" in held_rows["대진초 집중돌봄교실 외부 강사 모집"][1]
    assert held_rows["청소년방과후아카데미 기간제근로자 채용"][0] == "보류"
    assert "판별 유보" in held_rows["청소년방과후아카데미 기간제근로자 채용"][1]
    assert ws.max_row == 4  # 마감 지난 공고는 싣지 않는다


def test_build_refuses_changed_template(tmp_path):
    wb = load_workbook(TEMPLATE)
    wb["공고"]["D1"] = "접수마감"
    changed = tmp_path / "changed.xlsx"
    wb.save(changed)
    with pytest.raises(ValueError, match="칸 이름"):
        build_gangsaitda(_store(tmp_path), NOW, template=changed)


# ---- 본문 저장·백필 ---------------------------------------------------

def test_merge_carries_body_excerpt():
    from gia.dedupe import merge
    old = make_posting("요가 강사 모집")
    new = make_posting("요가 강사 모집", body_excerpt="1. 계약기간: 2026. 11. 1. ~ 12. 31.")
    merged, changed = merge(old, new, NOW)
    assert merged.body_excerpt.startswith("1. 계약기간")
    assert changed is False  # 본문을 채운 것만으로 '변경 공고' 알림을 내지 않는다


def test_content_selector_narrows_stored_excerpt(settings):
    from gia.collectors.base import HttpClient
    from gia.collectors.html_list import HtmlListAdapter
    from gia.models import RawListing
    from gia.pipeline import body_excerpt

    page = ("<html><body><div class='v'><dl><dt>전화번호</dt><dd>055-000-0000</dd></dl>"
            "<div class='txt'><p>1. 계약기간: 2026. 11. 1. ~ 12. 31.</p><p>접수: a@school.kr</p></div></div></body></html>")
    http = HttpClient(settings.collector, transport=httpx.MockTransport(
        lambda r: httpx.Response(200, text=page, headers={"content-type": "text/html; charset=utf-8"})))
    cfg = make_source("gne", type="html_list", list_url="https://x.org/list", row_selector="tr",
                      detail={"fetch": True, "body_selector": "div.v", "content_selector": "div.v div.txt"})
    raw = HtmlListAdapter(cfg, http, settings.collector).fetch_detail(
        RawListing(source_id="gne", title="t", url="https://x.org/view/1"))
    assert "055-000-0000" in raw.body_text  # 판별용 본문은 머리 표까지
    excerpt = body_excerpt(raw)
    assert excerpt.startswith("1. 계약기간") and "055" not in excerpt
    assert "[이메일]" in excerpt and "a@school.kr" not in excerpt  # 공개 저장소에 들어가므로 가린다
    http.close()


def test_backfill_fills_old_postings_by_url(tmp_path, settings):
    from gia.pipeline import SourceOutcome, backfill_bodies
    from gia.models import SourceRunResult

    old = make_posting("요가 강사 모집", status=Status.active,
                       sources=[SourceRef(source_id="gne", url="https://x.org/view/9", fetched_at=NOW)])
    other = make_posting("다른 소스 공고", status=Status.active)  # source_id='src' — 건드리지 않는다
    gone = make_posting("마감 공고", status=Status.expired,
                        sources=[SourceRef(source_id="gne", url="https://x.org/view/8", fetched_at=NOW)])
    store = _store(tmp_path, old, other, gone)

    class Stub:
        calls: list[str] = []

        def fetch_detail(self, listing):
            self.calls.append(listing.url)
            return RawPosting(**listing.model_dump(), body_text="2. 계약기간: 2026. 11. 1. ~ 12. 31.\n접수: a@b.kr",
                              fetched_at=NOW)

    stub = Stub()
    out = SourceOutcome(result=SourceRunResult(source_id="gne"))
    backfill_bodies(make_source("gne"), stub, store, out, None, settings.collector, t0=__import__("time").monotonic())
    assert stub.calls == ["https://x.org/view/9"]
    assert store.postings[old.canonical_key].body_excerpt.startswith("2. 계약기간")
    assert store.postings[other.canonical_key].body_excerpt == ""
    assert out.result.detail_fetched == 1

    # 한 번 채운 공고는 다시 열지 않는다
    stub.calls.clear()
    backfill_bodies(make_source("gne"), stub, store, out, None, settings.collector, t0=__import__("time").monotonic())
    assert stub.calls == []


def test_report_email_attaches_gangsaitda_file(tmp_path, monkeypatch):
    """--send 경로에서 요약 엑셀과 강사잇다 양식이 함께 붙고, 본문에 줄 수가 적힌다."""
    import gia.notify.email as email_mod
    from gia import __main__ as cli

    sent: dict = {}

    def fake_send(cfg, subject, html, text, smtp_factory=None, attachments=None):
        sent["names"] = [n for n, _ in attachments or []]
        sent["html"] = html
        return 1

    monkeypatch.setattr(email_mod, "send_email", fake_send)
    for k, v in {"SMTP_HOST": "smtp.example.org", "EMAIL_TO": "me@example.org"}.items():
        monkeypatch.setenv(k, v)
    import shutil
    from pathlib import Path
    repo = Path(cli.__file__).resolve().parents[1]
    work = tmp_path / "work"
    shutil.copytree(repo / "config", work / "config")
    (work / "data").mkdir()
    settings_file = work / "config" / "settings.yaml"
    text = settings_file.read_text(encoding="utf-8")
    import re as _re
    settings_file.write_text(_re.sub(r"channels:\s*\[[^\]]*\]", "channels: [email]", text), encoding="utf-8")
    monkeypatch.chdir(work)

    rc = cli.main(["report", "--send", "--no-mark"])
    assert rc == 0
    assert any(n.startswith("경남_강사구인공고_") for n in sent["names"])
    assert any(n.startswith("강사잇다_공고_") for n in sent["names"])
    assert "강사잇다 양식 첨부" in sent["html"]


def test_excerpt_keeps_attachment_text_when_content_selector_used():
    """본문 칸은 제목 한 줄이고 내용이 HWP 첨부에 있는 곳(창원시설공단): 발췌에 첨부 글을 붙인다."""
    from gia.pipeline import body_excerpt
    raw = RawPosting(source_id="cw_fmc", title="t", url="https://x.org/1", fetched_at=NOW,
                     body_text="작성자 홍길동\n요가 강사 모집\n\n[첨부: 공고.hwp]\n1. 계약기간: 2026. 11. 1. ~ 12. 31.",
                     extra={"content_text": "요가 강사 모집",
                            "attachment_text": "[첨부: 공고.hwp]\n1. 계약기간: 2026. 11. 1. ~ 12. 31."})
    ex = body_excerpt(raw)
    assert ex.startswith("요가 강사 모집")
    assert "1. 계약기간" in ex
    assert "홍길동" not in ex  # 작성자 실명은 본문 칸 밖이라 저장하지 않는다


def test_empty_content_does_not_fall_back_to_page_chrome():
    """본문 칸이 비어도 작성자 실명이 있는 전체 틀로 물러나지 않는다(늘푸른전당, 2026-09-28)."""
    from gia.pipeline import body_excerpt
    raw = RawPosting(source_id="cw_fmc", title="t", url="https://x.org/1", fetched_at=NOW,
                     body_text="작성자\n정민경\n본문\n이전글\n다른 공고", extra={"content_text": ""})
    assert body_excerpt(raw) == ""


def test_clean_detail_drops_attachment_markers():
    assert clean_detail("[첨부: %EC%9D%91.hwp]\n1. 모집분야: 요가") == "1. 모집분야: 요가"
