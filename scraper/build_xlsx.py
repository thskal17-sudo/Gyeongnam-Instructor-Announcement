#!/usr/bin/env python3
"""data/announcements.json + data/institutions.json -> data/announcements.xlsx"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "announcements.xlsx"
FONT = "맑은 고딕"
REGION_ORDER = {"부산": 0, "울산": 1, "경남": 2}
CATEGORIES = ["강사모집", "강좌개설신청", "강좌개설공모", "강좌개설제안(상시)", "강사인증/등록", "강사공지", "일반강사채용(참고)"]

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name=FONT, size=10)
LINK_FONT = Font(name=FONT, size=10, color="0563C1", underline="single")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_table(ws, headers, rows, widths, name, link_col=None):
    ws.append(headers)
    style_header(ws, 1, len(headers))
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.font, cell.border = BODY_FONT, BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if link_col and row[link_col - 1].value:
            row[link_col - 1].hyperlink = row[link_col - 1].value
            row[link_col - 1].font = LINK_FONT
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 2)}"


def build() -> None:
    ann = json.loads((ROOT / "data" / "announcements.json").read_text(encoding="utf-8"))
    inst = json.loads((ROOT / "data" / "institutions.json").read_text(encoding="utf-8"))
    ann.sort(key=lambda r: (REGION_ORDER.get(r["region"], 9), r["university"], r.get("posted_date") or "", r["title"]))
    inst.sort(key=lambda r: (REGION_ORDER.get(r["region"], 9), r["university"]))

    wb = Workbook()

    # ---- 시트 1: 공고목록 ----
    ws = wb.active
    ws.title = "공고목록"
    headers = ["번호", "지역", "대학", "기관", "구분", "공고 제목", "게시일", "링크", "요약", "게시판", "출처", "신뢰도", "수집일"]
    rows = [[i, a["region"], a["university"], a["center"], a["category"], a["title"], a.get("posted_date") or "",
             a["url"], a["summary"], a.get("board", ""), a["source"], a["confidence"], a["collected_at"]]
            for i, a in enumerate(ann, 1)]
    write_table(ws, headers, rows, [6, 7, 16, 18, 16, 48, 11, 40, 60, 14, 12, 8, 11], "공고", link_col=8)

    # ---- 시트 2: 기관목록 ----
    ws2 = wb.create_sheet("기관목록")
    headers2 = ["지역", "대학", "기관", "홈페이지", "공고 게시판", "전화", "이메일", "수집 공고 수(참고 제외)"]
    rows2 = []
    for k, i in enumerate(inst, 2):
        boards = "\n".join(f"{b['name']}: {b['url']}" for b in i.get("boards", [])) or "(게시판 URL 미확인)"
        rows2.append([i["region"], i["university"], i["center"], i["homepage"], boards, i.get("phone", ""), i.get("email", ""),
                      f'=COUNTIFS(공고목록!$C:$C,B{k},공고목록!$E:$E,"<>일반강사채용(참고)")'])
    write_table(ws2, headers2, rows2, [7, 20, 22, 36, 60, 20, 26, 12], "기관", link_col=4)

    # ---- 시트 3: 요약 (COUNTIFS 로 공고목록에서 집계) ----
    ws3 = wb.create_sheet("요약")
    ws3["A1"] = "지역 × 구분별 공고 건수 (공고목록 시트 기준 자동 집계)"
    ws3["A1"].font = Font(name=FONT, bold=True, size=12)
    ws3.append([])
    ws3.append(["지역"] + CATEGORIES + ["합계"])
    style_header(ws3, 3, len(CATEGORIES) + 2)
    regions = ["부산", "울산", "경남"]
    for ri, region in enumerate(regions, 4):
        ws3.cell(row=ri, column=1, value=region)
        for ci, cat in enumerate(CATEGORIES, 2):
            col = get_column_letter(ci)
            ws3.cell(row=ri, column=ci, value=f'=COUNTIFS(공고목록!$B:$B,$A{ri},공고목록!$E:$E,{col}$3)')
        last = get_column_letter(len(CATEGORIES) + 1)
        ws3.cell(row=ri, column=len(CATEGORIES) + 2, value=f"=SUM(B{ri}:{last}{ri})")
    tr = 4 + len(regions)
    ws3.cell(row=tr, column=1, value="합계")
    for ci in range(2, len(CATEGORIES) + 3):
        col = get_column_letter(ci)
        ws3.cell(row=tr, column=ci, value=f"=SUM({col}4:{col}{tr - 1})")
    for row in ws3.iter_rows(min_row=4, max_row=tr, max_col=len(CATEGORIES) + 2):
        for cell in row:
            cell.font, cell.border = BODY_FONT, BORDER
            cell.alignment = Alignment(horizontal="center")
    for cell in ws3[tr]:
        cell.font = Font(name=FONT, bold=True, size=10)
    ws3.column_dimensions["A"].width = 10
    for ci in range(2, len(CATEGORIES) + 3):
        ws3.column_dimensions[get_column_letter(ci)].width = 17
    ws3.row_dimensions[3].height = 30

    note = tr + 2
    notes = [
        "구분 설명",
        "강사모집: 강사를 직접 공모하는 공고 / 강좌개설신청·공모: 강사가 강좌 개설을 신청하는 방식의 학기별 공고",
        "강좌개설제안(상시): 공고 없이 홈페이지 메뉴로 상시 접수 / 강사인증/등록: 강사 풀 등록 / 강사공지: 기존 강사 대상 안내",
        "일반강사채용(참고): 평생교육원이 아닌 대학 본부의 학부 강사·겸임교수 채용 (참고용)",
        "신뢰도: high = 게시글 URL까지 확인, medium = 제목·게시판만 확인, low = 검색 스니펫으로만 확인",
        "출처 web_search: 대학 도메인 직접 접속이 차단된 환경에서 웹 검색 결과로 수집 / scraper: 자동 수집",
    ]
    for j, t in enumerate(notes):
        c = ws3.cell(row=note + j, column=1, value=t)
        c.font = Font(name=FONT, size=9, bold=(j == 0), color="595959")

    wb.calculation.fullCalcOnLoad = True  # 캐시값이 없어도 Excel 이 열 때 전체 재계산
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)} ({len(ann)} announcements, {len(inst)} institutions)")


if __name__ == "__main__":
    build()
