"""강사잇다 '공고 올리기' 양식(.xlsx)을 채운다.

양식 파일은 사이트가 준 것을 그대로 쓴다(config/templates/). '공고' 시트에 공고를 한 줄씩
적고 '예시'·'안내' 시트는 건드리지 않는다. 칸 이름이 바뀌었으면 채우지 않고 멈춘다 —
엉뚱한 칸에 들어간 파일이 사이트에 올라가는 쪽이 더 나쁘다.

필수 칸(제목·기관명·지역·마감일·수업 일정·상세 내용)을 다 못 채운 줄도 지우지 않고,
양식 안내가 허용하는 '처리=보류'로 남긴다. 사이트는 보류 줄을 올리지 않으므로 원문을 보고
빈칸을 채운 뒤 '보류'만 지우면 된다. 빈칸을 그럴듯한 말로 메우지 않는다.
"""
from __future__ import annotations

from copy import copy
from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment

from ..extract.deadline import KST
from ..extract.detail import (
    clean_detail, clean_title, expand_school_name, extract_documents, extract_headcount,
    extract_qualifications, extract_schedule, extract_target,
)
from ..models import OrgType, Posting, Status
from ..normalize import SIGUN_SUFFIX
from ..store import Store

TEMPLATE = Path(__file__).resolve().parents[2] / "config" / "templates" / "gangsaitda-job-template.xlsx"
SHEET = "공고"
COLS = ["제목", "기관명", "지역", "마감일", "수업 일정", "상세 내용",
        "수업 대상", "모집 인원", "지원 자격", "제출 서류", "원문 링크", "지원서 링크", "지원 이메일"]
REQUIRED = COLS[:6]
EXTRA = ["처리", "메모"]  # 양식 안내 10번: '처리' 칸에 '보류'로 적은 줄은 올리지 않고, '메모' 같은 다른 칸은 무시한다


def region_label(regions: list[str]) -> str:
    """['창원'] → '경남 창원시'. 양식 예시가 '부산 해운대구' 꼴이다."""
    for r in regions or []:
        if r in SIGUN_SUFFIX:
            return f"경남 {r}{SIGUN_SUFFIX[r]}"
    return "경남"


def to_row(p: Posting) -> dict[str, object]:
    body = p.body_excerpt or ""
    title = clean_title(p.title)
    url = p.primary_url if p.primary_url.startswith("https://") else ""  # 양식: https:// 로 시작
    return {
        "제목": title,
        "기관명": expand_school_name(p.org_name) if p.org_type == OrgType.edu_office else p.org_name,
        "지역": region_label(p.region),
        "마감일": p.deadline.astimezone(KST).date().isoformat() if p.deadline else "",
        # 접수 마감 뒤에 시작하는 날짜 범위만 일정으로 본다 — 표 형식 공고에서 접수 기간을 거르려고
        "수업 일정": extract_schedule(body, not_before=p.deadline.astimezone(KST).date() if p.deadline else None),
        "상세 내용": clean_detail(body, title),
        "수업 대상": extract_target(body),
        "모집 인원": extract_headcount(body),
        "지원 자격": "\n".join(p.qualifications or extract_qualifications(body)),
        "제출 서류": extract_documents(body),
        "원문 링크": url,
        "지원서 링크": "",
        # 접수 메일은 채우지 않는다. 원문에서 가져오려면 가리지 않은 주소를 공개 저장소(data/)에
        # 쌓아야 하는데, 교사 개인 계정이 많아 수집 봇의 표적이 된다. 원문 링크로 확인할 수 있다
        "지원 이메일": "",
    }


def hold_reason(p: Posting, row: dict[str, object]) -> str:
    """보류로 둘 이유. 없으면 빈 문자열(바로 올릴 수 있는 줄)."""
    why: list[str] = []
    missing = [c for c in REQUIRED if not row.get(c)]
    if missing:
        why.append("빈 필수 칸: " + ", ".join(missing))
    if "판별유보" in p.flags:
        why.append(f"강사 공고인지 판별 유보(점수 {p.relevance_score})")
    if not why:
        return ""
    return " · ".join(why) + " — 원문을 확인해 채운 뒤 '보류'를 지우세요"


def select_postings(store: Store, now: datetime) -> list[Posting]:
    """올릴 후보: 진행 중이고 마감이 지나지 않은 공고 전부.

    매일 신규만이 아니라 전체를 보낸다. 사이트가 같은 제목은 건너뛰므로(양식 안내 5번)
    어제 못 올린 줄도 다시 실린다.
    """
    return [p for p in store.values()
            if p.status != Status.expired
            and not (p.deadline and p.deadline < now)
            and "피드백제외" not in p.flags]


def build_gangsaitda(store: Store, now: datetime, template: Path = TEMPLATE) -> tuple[bytes, int, int]:
    """(파일 바이트, 바로 올릴 줄 수, 보류 줄 수)."""
    wb = load_workbook(template)
    ws = wb[SHEET]
    header = [ws.cell(1, i + 1).value for i in range(len(COLS))]
    if header != COLS:
        raise ValueError(f"강사잇다 양식 칸 이름이 바뀌었습니다: {header}")

    ready: list[tuple[Posting, dict]] = []
    held: list[tuple[Posting, dict, str]] = []
    for p in select_postings(store, now):
        row = to_row(p)
        why = hold_reason(p, row)
        if why:
            held.append((p, row, why))
        else:
            ready.append((p, row))
    key = lambda item: (item[0].deadline is None, item[0].deadline or now)  # noqa: E731
    ready.sort(key=key)
    held.sort(key=key)

    # 처리·메모 머리는 선택 칸(회색) 모양을 빌린다
    style_src = ws.cell(1, COLS.index("수업 대상") + 1)
    for j, name in enumerate(EXTRA):
        c = ws.cell(1, len(COLS) + 1 + j, name)
        c.font, c.fill, c.alignment, c.border = copy(style_src.font), copy(style_src.fill), copy(style_src.alignment), copy(style_src.border)
    ws.column_dimensions["N"].width = 8
    ws.column_dimensions["O"].width = 48

    wrap = Alignment(wrap_text=True, vertical="top")  # '예시' 시트와 같은 모양
    r = 2
    for p, row, *rest in [*ready, *held]:
        values = [row[c] for c in COLS] + (["보류", rest[0]] if rest else ["", ""])
        for i, v in enumerate(values, start=1):
            cell = ws.cell(r, i, v if v not in (None, "") else None)
            cell.alignment = wrap
        r += 1

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), len(ready), len(held)


def gangsaitda_name(now: datetime) -> str:
    return f"강사잇다_공고_{now.astimezone(KST):%Y-%m-%d}.xlsx"
